"""
services/report_service.py - Report Generation & Export

Uses Pandas for data processing and openpyxl for Excel generation.
Supports:
  - Daily reports
  - Monthly reports
  - Student-wise reports
  - CSV export
  - Excel export (.xlsx)
"""

import os
import io
from datetime import date, datetime

import pandas as pd
from flask import send_file

from utils.database import db
from models.student import Student
from models.attendance import Attendance
from config import ATTENDANCE_STATUS_PRESENT


# ─── Build DataFrames ─────────────────────────────────────────────────────────

def build_attendance_dataframe(
    date_filter=None,
    student_id_filter=None,
    dept_filter=None,
    month=None,
    year=None
) -> pd.DataFrame:
    """
    Query attendance records and return as a Pandas DataFrame.
    All parameters are optional filters.
    """
    query = (
        db.session.query(
            Student.student_id.label('Student ID'),
            Student.name.label('Name'),
            Student.department.label('Department'),
            Student.year.label('Year'),
            Attendance.date.label('Date'),
            Attendance.time.label('Time'),
            Attendance.status.label('Status'),
            Attendance.confidence.label('Confidence (%)'),
        )
        .join(Student, Attendance.student_id == Student.id)
    )

    if date_filter:
        query = query.filter(Attendance.date == date_filter)
    if student_id_filter:
        query = query.filter(Student.student_id == student_id_filter)
    if dept_filter:
        query = query.filter(Student.department == dept_filter)
    if month and year:
        month_str = f'{year}-{int(month):02d}'
        query = query.filter(Attendance.date.like(f'{month_str}-%'))

    rows = query.order_by(Attendance.date.desc(), Attendance.time.desc()).all()

    if not rows:
        return pd.DataFrame(columns=[
            'Student ID', 'Name', 'Department', 'Year',
            'Date', 'Time', 'Status', 'Confidence (%)'
        ])

    df = pd.DataFrame(rows, columns=[
        'Student ID', 'Name', 'Department', 'Year',
        'Date', 'Time', 'Status', 'Confidence (%)'
    ])
    return df


def build_daily_summary_df(date_str: str) -> pd.DataFrame:
    """
    Build a daily summary showing every student and their attendance status
    (Present / Absent) for a specific date.
    """
    all_students = Student.query.order_by(Student.name).all()

    # Get the set of student IDs that are present on this date
    present_ids = {
        r.student_id
        for r in Attendance.query.filter_by(date=date_str, status=ATTENDANCE_STATUS_PRESENT).all()
    }

    rows = []
    for s in all_students:
        status = ATTENDANCE_STATUS_PRESENT if s.id in present_ids else 'Absent'
        time   = ''
        if s.id in present_ids:
            rec  = Attendance.query.filter_by(student_id=s.id, date=date_str).first()
            time = rec.time if rec else ''
        rows.append({
            'Student ID': s.student_id,
            'Name':       s.name,
            'Department': s.department,
            'Year':       s.year,
            'Status':     status,
            'Time':       time
        })

    return pd.DataFrame(rows)


def build_monthly_summary_df(year: int, month: int) -> pd.DataFrame:
    """
    Build a monthly summary with per-student present/absent/percentage.
    """
    import calendar
    _, days_in_month = calendar.monthrange(year, month)
    month_str = f'{year}-{month:02d}'

    all_students = Student.query.order_by(Student.name).all()
    rows = []

    for s in all_students:
        present = Attendance.query.filter(
            Attendance.student_id == s.id,
            Attendance.date.like(f'{month_str}-%'),
            Attendance.status == ATTENDANCE_STATUS_PRESENT
        ).count()
        absent = days_in_month - present
        pct    = round((present / days_in_month * 100), 2)
        rows.append({
            'Student ID':  s.student_id,
            'Name':        s.name,
            'Department':  s.department,
            'Year':        s.year,
            'Present Days': present,
            'Absent Days':  absent,
            'Total Days':   days_in_month,
            'Percentage (%)': pct
        })

    return pd.DataFrame(rows)


# ─── Export: CSV ──────────────────────────────────────────────────────────────

def export_csv(df: pd.DataFrame, filename: str):
    """
    Return a CSV file as a Flask response for download.
    """
    output = io.BytesIO()
    df.to_csv(output, index=False, encoding='utf-8-sig')   # utf-8-sig for Excel compatibility
    output.seek(0)
    return send_file(
        output,
        mimetype='text/csv',
        as_attachment=True,
        download_name=filename
    )


# ─── Export: Excel ────────────────────────────────────────────────────────────

def export_excel(df: pd.DataFrame, filename: str, sheet_name: str = 'Attendance'):
    """
    Return an Excel (.xlsx) file as a Flask response for download.
    Uses openpyxl engine for styling.
    """
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)

        # ── Basic styling ──────────────────────────────────────────────────────
        workbook  = writer.book
        worksheet = writer.sheets[sheet_name]

        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
        center_align = Alignment(horizontal='center', vertical='center')
        thin_border  = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # Style header row
        for col_num, cell in enumerate(worksheet[1], 1):
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = center_align
            cell.border    = thin_border

        # Auto-fit columns
        for col_num, col in enumerate(worksheet.columns, 1):
            max_len = max((len(str(c.value or '')) for c in col), default=10)
            col_letter = get_column_letter(col_num)
            worksheet.column_dimensions[col_letter].width = min(max_len + 4, 40)

        # Alternate row colours
        light_fill = PatternFill(start_color='EEF4FF', end_color='EEF4FF', fill_type='solid')
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=2), 2):
            for cell in row:
                cell.border    = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if row_idx % 2 == 0:
                    cell.fill = light_fill

        worksheet.freeze_panes = 'A2'   # Freeze header row

    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


# ─── Filename Helpers ─────────────────────────────────────────────────────────

def daily_filename(date_str: str, ext: str) -> str:
    return f'attendance_{date_str}.{ext}'


def monthly_filename(year: int, month: int, ext: str) -> str:
    return f'attendance_{year}-{month:02d}.{ext}'


def student_filename(student_id: str, ext: str) -> str:
    return f'student_{student_id}_attendance.{ext}'
